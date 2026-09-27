#!/usr/bin/env python3
"""Publish the public map and private BMA needs from the live Google workbook."""

from __future__ import annotations

from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import logging
import math
import os
from pathlib import Path
import time
from uuid import uuid4

import openpyxl
import requests

from sync import POI_CACHE, POI_URL, STATE_DIR, WEBROOT, cached_point, cache_expired, read_json, utc_now, valid_point, write_json_atomic


SHEET_ID = "1S3mDcZisdLkcIFbv996xKo1G6-E2KVEt1d8lyz-gR_U"
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit"
EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
PUBLIC_OUTPUT = WEBROOT / "floodsupport-data.json"
STAFF_OUTPUT = STATE_DIR / "staff-data.json"
RAW_OUTPUT = STATE_DIR / "source-latest.xlsx"
COORD_SEED = STATE_DIR / "coordinate-seed.json"
LOG = STATE_DIR / "sheet-sync.log"
SHELTER = "ศูนย์พักพิงชั่วคราว"
PARKING = "จุดจอดรถ"


def text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def number(value: object) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        n = float(str(value).replace(",", "").strip())
        return int(n) if math.isfinite(n) and n >= 0 and n.is_integer() else None
    except ValueError:
        return None


def value(row: tuple, index: int) -> object:
    return row[index] if index < len(row) else None


def district_name(value_: object, valid: set[str]) -> str | None:
    raw = text(value_)
    name = raw.removeprefix("เขต").strip()
    return name if raw.startswith("เขต") and name in valid else None


def stable_id(kind: str, district: str, subdistrict: str, name: str) -> str:
    return sha256("\x1f".join((kind, district, subdistrict, name)).encode("utf-8")).hexdigest()[:24]


def download_workbook() -> tuple[object, bytes]:
    error = None
    for attempt in range(2):
        try:
            response = requests.get(EXPORT_URL, timeout=(8, 30), headers={"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
            response.raise_for_status()
            if not 100_000 <= len(response.content) <= 15_000_000 or not response.content.startswith(b"PK"):
                raise ValueError("Workbook download has an unexpected format or size")
            workbook = openpyxl.load_workbook(BytesIO(response.content), read_only=True, data_only=True)
            return workbook, response.content
        except (requests.RequestException, ValueError, OSError, openpyxl.utils.exceptions.InvalidFileException) as exc:
            error = exc
            logging.warning("Google Sheets download attempt %s failed: %s", attempt + 1, exc)
            if attempt == 0:
                time.sleep(2)
    raise RuntimeError("Google Sheets workbook could not be downloaded") from error


def sheet(workbook: object, prefix: str, checks: list[tuple[int, int, str]]):
    matches = [item for item in workbook.worksheets if item.title.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one worksheet beginning with {prefix!r}")
    result = matches[0]
    for row, col, expected in checks:
        if expected not in text(result.cell(row, col).value):
            raise ValueError(f"Unexpected header in {result.title}!{result.cell(row, col).coordinate}")
    return result


def query_poi(name: str, district: str) -> dict | None:
    escaped_name = name.replace("'", "''")
    escaped_district = district.replace("'", "''")
    response = requests.get(POI_URL, params={
        "where": f"NAME='{escaped_name}' AND DISTRICT='{escaped_district}'",
        "outFields": "OBJECTID,NAME,DISTRICT", "returnGeometry": "true", "outSR": "4326", "f": "json",
    }, timeout=(5, 8))
    response.raise_for_status()
    payload = response.json()
    if payload.get("error"):
        raise ValueError(str(payload["error"]))
    features = payload.get("features")
    if not isinstance(features, list):
        raise ValueError("POI response has no features")
    if len(features) != 1:
        return None
    feature = features[0]
    attrs, geometry = feature.get("attributes") or {}, feature.get("geometry") or {}
    lat, lon = geometry.get("y"), geometry.get("x")
    if attrs.get("NAME") != name or attrs.get("DISTRICT") != district or not valid_point(lat, lon):
        return None
    return {"lat": float(lat), "lon": float(lon), "method": "bma-poi-exact", "poiObjectId": attrs.get("OBJECTID")}


def coordinate_map(places: list[dict], all_names: Counter) -> dict[tuple[str, str], dict]:
    cache = read_json(POI_CACHE, {})
    seed = read_json(COORD_SEED, {})
    if not isinstance(cache, dict) or not isinstance(seed, dict):
        raise ValueError("Coordinate cache has an invalid format")
    unique = {(row["name"], row["district"]) for row in places if all_names[(row["category"], row["name"], row["district"])] == 1}
    missing = sorted((name, district) for name, district in unique if cache_expired(cache.get(json.dumps([name, district], ensure_ascii=False))))[:24]
    if missing:
        with ThreadPoolExecutor(max_workers=6) as pool:
            futures = {pool.submit(query_poi, name, district): (name, district) for name, district in missing}
            for future in as_completed(futures):
                name, district = futures[future]
                try:
                    cache[json.dumps([name, district], ensure_ascii=False)] = {"checkedAt": utc_now(), "point": future.result()}
                except Exception as exc:
                    logging.warning("POI lookup failed for %r, %r: %s", name, district, exc)
        write_json_atomic(POI_CACHE, cache)
    points = {}
    for name, district in unique:
        key = json.dumps([name, district], ensure_ascii=False)
        poi = cached_point(cache.get(key))
        old = seed.get(key)
        seeded = old if isinstance(old, dict) and valid_point(old.get("lat"), old.get("lon")) else None
        if poi and seeded:
            # A name reused for a different place must never inherit the old pin.
            la, lb = math.radians(poi["lat"]), math.radians(seeded["lat"])
            h = math.sin((lb - la) / 2) ** 2 + math.cos(la) * math.cos(lb) * math.sin(math.radians(seeded["lon"] - poi["lon"]) / 2) ** 2
            if 6371000 * 2 * math.asin(min(1, math.sqrt(h))) > 1500:
                continue
        point = poi or seeded
        if point:
            points[(name, district)] = point
    return points


def write_bytes_atomic(path: Path, contents: bytes) -> None:
    temporary = path.with_name(path.name + "." + uuid4().hex + ".tmp")
    try:
        temporary.write_bytes(contents)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
