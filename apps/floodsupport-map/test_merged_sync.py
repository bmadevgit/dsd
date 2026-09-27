import json
import unittest
from unittest.mock import patch
from openpyxl import Workbook

import merged_sync
import sheet_parser


class MergeTests(unittest.TestCase):
    def test_duplicate_sheet_site_keeps_populated_row_but_rejects_conflict(self):
        complete = {"id": "site", "capacity": 60, "occupied": 40, "available": 20, "location": None}
        blank = {"id": "site", "capacity": None, "occupied": None, "available": None, "location": None}
        rows = sheet_parser.deduplicate_places([complete, blank])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["capacity"], 60)
        self.assertEqual(rows[0]["duplicateSourceRows"], 2)
        with self.assertRaises(ValueError):
            sheet_parser.deduplicate_places([complete, {**blank, "capacity": 70}])

    def test_grouped_water_header_uses_remaining_support_column(self):
        book = Workbook()
        sheet = book.active
        sheet.cell(1, 8, "น้ำดื่ม")
        sheet.cell(1, 12, "รถ GMC")
        sheet.cell(4, 8, "ต้องการสนับสนุน (ขาด)")
        sheet.cell(4, 11, "คงเหลือ สนับสนุนเพิ่มเติม (ขวด)")
        self.assertEqual(sheet_parser.need_col(sheet, "น้ำดื่ม"), 10)

    def test_api_wins_fields_and_sheet_supplies_coordinate(self):
        api_row = {"id": "a", "name": "ศูนย์ ก", "district": "คลองเตย", "category": "ศูนย์พักพิงชั่วคราว",
                   "status": "เต็ม", "capacity": 50, "occupied": 50, "available": 0, "link": None}
        sheet_row = {"id": "s", "name": "ศูนย์ ก", "district": "คลองเตย", "category": "ศูนย์พักพิงชั่วคราว",
                     "status": "เปิดให้บริการ", "capacity": 80, "available": 40, "phone": "123", "conditions": "ฟรี",
                     "location": {"lat": 13.75, "lon": 100.5, "method": "google-sheet-coordinate"}}
        with patch.object(merged_sync, "read_json", return_value={}):
            result, staff = merged_sync.merge(
                {"lastFetchedAt": merged_sync.utc_now(), "payload": {"facilities": [api_row], "categories": ["ศูนย์พักพิงชั่วคราว"]}},
                {"lastFetchedAt": merged_sync.utc_now(), "places": [sheet_row], "staff": {
                    "districtNeeds": [{"district": "คลองเตย", "needs": {"water": 25}}],
                    "shelterNeeds": [{"id": "a" * 24, "name": "ศูนย์ ก", "district": "คลองเตย",
                                      "request": "น้ำดื่ม", "capacity": 80, "occupied": 40, "available": 40}]}},
                False, False)
        self.assertEqual(result["total"], 1)
        row = result["facilities"][0]
        self.assertEqual((row["status"], row["capacity"]), ("เต็ม", 50))
        self.assertEqual(row["location"]["method"], "google-sheet-coordinate")
        self.assertEqual(row["sources"], ["Flood Support", "Google Sheets"])
        self.assertEqual(staff["districtNeeds"][0]["source"], "Google Sheets")
        self.assertEqual(staff["shelterNeeds"][0]["requestSource"], "Google Sheets")
        self.assertEqual(staff["shelterNeeds"][0]["floodSupport"]["capacity"], 50)

    def test_ambiguous_shelter_need_does_not_claim_api_match(self):
        api_row = {"id": "a", "name": "ศูนย์ ก", "district": "คลองเตย", "category": merged_sync.sheet_sync.SHELTER,
                   "status": "เปิด", "capacity": 50, "occupied": 0, "available": 50, "unitType": "คน", "link": None}
        sheet_rows = [{"id": str(n) * 24, "name": "ศูนย์ ก", "district": "คลองเตย", "request": "น้ำดื่ม"}
                      for n in (1, 2)]
        with patch.object(merged_sync, "read_json", return_value={}):
            _, staff = merged_sync.merge(
                {"lastFetchedAt": merged_sync.utc_now(), "payload": {"facilities": [api_row], "categories": [api_row["category"]]}},
                {"lastFetchedAt": merged_sync.utc_now(), "places": [], "staff": {"districtNeeds": [], "shelterNeeds": sheet_rows}},
                False, False)
        self.assertTrue(all("floodSupport" not in row for row in staff["shelterNeeds"]))

    def test_invalid_sheet_coordinate_has_no_pin(self):
        self.assertIsNone(sheet_parser.coordinates(("#N/A", "#N/A"), 0, 1))
        self.assertIsNone(sheet_parser.coordinates((15, 100.5), 0, 1))
        self.assertEqual(sheet_parser.coordinates((13.75, 100.5), 0, 1)["method"], "google-sheet-coordinate")

    def test_source_failure_uses_last_valid_cache(self):
        old = {"lastFetchedAt": merged_sync.utc_now(), "payload": {"facilities": [], "categories": []}}
        with patch.object(merged_sync.requests, "get", side_effect=TimeoutError), patch.object(merged_sync, "read_json", return_value=old):
            result, failed = merged_sync.read_api()
        self.assertTrue(failed)
        self.assertEqual(result, old)

    def test_sheet_failure_uses_last_valid_cache(self):
        old = {"lastFetchedAt": merged_sync.utc_now(), "places": [], "staff": {"districtNeeds": [], "shelterNeeds": []}}
        with patch.object(merged_sync.sheet_sync, "download_workbook", side_effect=TimeoutError), patch.object(merged_sync, "read_json", return_value=old):
            result, failed = merged_sync.read_sheets()
        self.assertTrue(failed)
        self.assertEqual(result, old)


if __name__ == "__main__":
    unittest.main()
